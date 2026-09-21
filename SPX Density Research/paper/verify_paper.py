"""Reconcile printed navigation with the rendered PDF and retained evidence."""
from pathlib import Path
import json,re,argparse,hashlib
import fitz
from docx import Document

HERE=Path(__file__).resolve().parent
parser=argparse.ArgumentParser()
parser.add_argument('pdf',type=Path)
parser.add_argument('--update-pages',action='store_true')
args=parser.parse_args()
pdf=fitz.open(args.pdf)
nav=json.loads((HERE/'navigation.json').read_text())
assets=json.loads((HERE/'assets.json').read_text())
toc=pdf.get_toc()
assert len(toc)==len(nav['headings']),(len(toc),len(nav['headings']))
assert toc[0][1]=='1 Introduction'
main_first=toc[0][2]-1

def roman(n):
    out=''
    for v,s in [(1000,'m'),(900,'cm'),(500,'d'),(400,'cd'),(100,'c'),(90,'xc'),(50,'l'),(40,'xl'),(10,'x'),(9,'ix'),(5,'v'),(4,'iv'),(1,'i')]:
        while n>=v:out+=s;n-=v
    return out

page_map={}
positions={}
for h,t in zip(nav['headings'],toc):
    assert h['title']==t[1],(h,t)
    page_map[h['key']]=str(t[2]-main_first)
    positions[h['key']]=t[2]
for title in nav['front']:
    found=[i for i in range(1,main_first) if re.search(r'(?m)^'+re.escape(title)+r'\s*$',pdf[i].get_text())]
    assert len(found)==1,(title,found)
    key='front_'+title.lower().replace(' ','_');page_map[key]=roman(found[0]);positions[key]=found[0]+1
for kind,tag in [('equations','Equation'),('figures','Figure'),('tables','Table')]:
    for n in assets[kind]:
        pat=r'(?m)^'+tag+' '+n+r'\. '
        found=[i for i in range(main_first,len(pdf)) if re.search(pat,pdf[i].get_text())]
        assert len(found)==1,(tag,n,found)
        page_map[tag.lower()+n]=str(found[0]-main_first+1);positions[tag.lower()+n]=found[0]+1
old=json.loads((HERE/'page_map.json').read_text()) if (HERE/'page_map.json').exists() else {}
changes={k:[old.get(k),v] for k,v in page_map.items() if old.get(k)!=v}
if args.update_pages:(HERE/'page_map.json').write_text(json.dumps(page_map,indent=2))
doc=Document(HERE/'SPX-Option-Implied-Density-Final-Paper.docx')
navlinks=doc.element.xpath('//w:hyperlink[@w:anchor]')
expected=[]
for link in navlinks:
    key=link.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor')
    if key in positions:
        displayed=link.xpath('./w:r/w:t')[-1].text
        if not args.update_pages:assert displayed==page_map[key],(key,displayed,page_map[key])
        expected.append(positions[key]-1)
actual=[]
for i in range(2,main_first):
    actual.extend(l['page'] for l in sorted(pdf[i].get_links(),key=lambda l:(round(l['from'].y0,1),l['from'].x0)) if l['kind']==1)
assert expected==actual,('Internal navigation mismatch',len(expected),len(actual))
preserved=json.loads((HERE/'preserved_inputs_sha256.json').read_text())
changed_inputs=[name for name,sha in preserved.items() if hashlib.sha256((HERE.parent/name).read_bytes()).hexdigest()!=sha]
assert not changed_inputs,changed_inputs
assert len(doc.inline_shapes)==33
assert len(doc.tables)==38
assert len(doc.element.xpath('//m:oMathPara'))==46
assert not re.search(r'\{\{|FLOAT(?:START|END|NEXT)|COVERTITLE|BODYBEGIN|NAV\|','\n'.join(p.text for p in doc.paragraphs))
text='\n'.join(p.get_text() for p in pdf)
assert '\ufffd' not in text
math_export_errors=[]
for i,p in enumerate(pdf):
    for b in p.get_text('dict')['blocks']:
        if 'lines' not in b:continue
        for l in b['lines']:
            for s in l['spans']:
                if s['color']==8388608:math_export_errors.append([i+1,s['text']])
assert not math_export_errors,math_export_errors
fig4=positions['figure4']-1
assert 'Figure 4 introduces independent' in pdf[fig4].get_text()
assert pdf[fig4].get_text().index('Figure 4.')<pdf[fig4].get_text().index('Figure 4 introduces independent')
blank=[i+1 for i,p in enumerate(pdf) if len(re.sub(r'OPTION-IMPLIED DENSITIES\s*\|\s*JOEL CERRAGA|\d+|\s+','',p.get_text()))<3 and len(p.get_images())==0]
overflow=[]
for i,p in enumerate(pdf):
    for b in p.get_text('dict')['blocks']:
        if 'lines' not in b:continue
        for l in b['lines']:
            for s in l['spans']:
                if s['bbox'][0]<20 or s['bbox'][2]>p.rect.width-20:overflow.append([i+1,s['text'],s['bbox']])
report={'physical_pages':len(pdf),'front_matter_pages':main_first-1,'main_pages':len(pdf)-main_first,'equations':46,'figures':33,'research_tables':36,'references':23,'native_math_objects':len(doc.element.xpath('//m:oMath')),'verified_internal_navigation_links':len(expected),'preserved_files':len(preserved),'changed_preserved_files':changed_inputs,'page_map_changes':changes,'blank_pages':blank,'possible_horizontal_overflow':overflow,'positions':positions,'sha256_docx':hashlib.sha256((HERE/'SPX-Option-Implied-Density-Final-Paper.docx').read_bytes()).hexdigest(),'sha256_pdf':hashlib.sha256(args.pdf.read_bytes()).hexdigest()}
(HERE/'verification.json').write_text(json.dumps(report,indent=2))
print(json.dumps({k:v for k,v in report.items() if k not in ['positions','page_map_changes','possible_horizontal_overflow']},indent=2))
print('Navigation changes',len(changes),'Horizontal flags',len(overflow))
