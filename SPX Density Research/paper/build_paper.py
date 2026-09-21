"""Build the integrated paper from reviewed prose and the saved research assets.

Requires pandoc, python-docx and Pillow. Layout QA/export uses LibreOffice;
page_map.json is populated from the rendered PDF by verify_paper.py.
This script never runs or changes the numerical experiments.
"""
from pathlib import Path
import re, json, subprocess, argparse, copy
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
ASSETS=json.loads((HERE/'assets.json').read_text())
PAGES=json.loads((HERE/'page_map.json').read_text()) if (HERE/'page_map.json').exists() else {}
TITLE='Recovering Option-Implied Risk-Neutral Densities from SPX Prices'
SUBTITLE='Constrained estimation, synthetic validation and empirical comparisons'
FILENAME='SPX-Option-Implied-Density-Final-Paper'
ABSTRACT='''This project investigates how European option prices can be used to recover risk-neutral terminal distributions, and what numerical consistency and price agreement establish about the recovered density. A controlled lognormal experiment first verifies call-price differentiation and demonstrates its sensitivity to quote errors. A finite histogram estimator then prices bin probabilities exactly, imposes non-negativity, unit mass and a forward-consistent mean, and links maturities through normalised call inequalities. Analytical within-bin checks identify calendar violations that boundary-only checks miss.

Synthetic evaluation examines support, resolution, smoothing, strike coverage and supplied-forward sensitivity using single-lognormal and two-component mixture benchmarks. Repeated experiments separate noise effects from missing strike information and distinguish Monte Carlo precision from density uncertainty. The results show that similar price errors can accompany materially different density and tail errors, and that selecting smoothing for interior price prediction need not improve density recovery.

The empirical application uses documented exchange bid and offer observations at three snapshots across two dates, with three SPXW expiries per snapshot. Discount factors and forwards are inferred from matched call–put pairs. A minimum-curvature estimator satisfies all 5,518 retained option intervals across the primary fits at the stated numerical tolerance. One held-out validation fold remains numerically unresolved. Linear programmes quantify the range of event probabilities compatible with the same finite density family, while dated comparisons distinguish fixed expiries from an explicitly assumed common-sixty-day mixture.

The evidence supports a conditional conclusion: probability validity, calendar consistency and in-sample price compatibility are necessary checks, but do not identify a unique empirical density or validate extrapolated tail probabilities. The work provides a reproducible estimation process, retains failed numerical attempts and separates imposed restrictions from measured recovery accuracy. The empirical findings remain conditional on the small sample, inferred carry, finite support and stated interpolation assumptions.'''

# Line breaking changes typography only; the scientific expressions are unchanged.
BREAKS={
11:r'\begin{aligned}\widehat C(K_i)&=\sum_j A_{ij}w_j,\\ A_{ij}&=\frac{D(T)}{2\Delta s_j}\left[(b_j-K_i)_+^2-(a_j-K_i)_+^2\right].\end{aligned}',
12:r'\begin{aligned}\min_{\mathbf w}\quad &\frac{1}{N}\sum_{i=1}^{N}\left(\frac{\sum_j A_{ij}w_j-y_i}{F}\right)^2\\ &+\frac{\lambda}{J-2}\sum_{j=1}^{J-2}(g_{j+2}-2g_{j+1}+g_j)^2,\\ &\text{subject to Equation (10).}\end{aligned}',
16:r'\begin{aligned}\widehat g_m(x)&=\sum_{j=1}^{J}\frac{w_{mj}}{\Delta u}\mathbf 1_{[u_{j-1},u_j)}(x),\\ \widehat f_Q(s;T_m)&=\frac{1}{F_m}\widehat g_m\left(\frac{s}{F_m}\right).\end{aligned}',
17:r'\begin{aligned}\min_{\mathbf w_1,\ldots,\mathbf w_M}\quad &\sum_{m=1}^{M}\ell_m(\mathbf w_m)\\ \text{subject to}\quad &w_{mj}\geq0,\quad \mathbf1^\top\mathbf w_m=1,\quad\bar{\mathbf u}^{\top}\mathbf w_m=1,\\ &\mathbf B(\kappa)(\mathbf w_{m+1}-\mathbf w_m)\geq0,\\ &\kappa\in\mathcal K_m,\quad m=1,\ldots,M-1.\end{aligned}',
18:r'\begin{aligned}\Delta_m(u_{j-1}+z)&=\Delta_m(u_{j-1})-z\sum_{\ell=j}^{J}\delta_{m\ell}+\frac{\delta_{mj}}{2\Delta u}z^2,\\ &0\leq z\leq\Delta u.\end{aligned}',
20:r'\begin{aligned}g_{\mathrm{mix}}(x,T)&=\sum_{\ell=1}^{L}\pi_\ell g_{\mathrm{LN}}(x,T;\sigma_\ell),\\ C_{\mathrm{mix}}(K,T)&=\sum_{\ell=1}^{L}\pi_\ell C_{\mathrm{BS}}(K,T;\sigma_\ell),\\ \sum_{\ell=1}^{L}\pi_\ell&=1.\end{aligned}',
21:r'\begin{aligned}R_{\Delta u}(\mathbf g)&=\Delta u\sum_{j=1}^{J-2}\left(\frac{g_{j+2}-2g_{j+1}+g_j}{(\Delta u)^2}\right)^2\\ &=\frac{1}{(\Delta u)^3}\sum_{j=1}^{J-2}(g_{j+2}-2g_{j+1}+g_j)^2.\end{aligned}',
23:r'\begin{aligned}L_1^{\mathrm{full}}(T)&=\int_{u_0}^{u_J}|\widehat g(x,T)-g(x,T)|\,dx\\ &\quad+Q(X_T<u_0)+Q(X_T>u_J),\\ X_T&=S_T/F(T).\end{aligned}',
24:r'\begin{aligned}p_{0.8}^{-}(T)&=Q(S_T<0.8F(T)),\\p_{1.2}^{+}(T)&=Q(S_T>1.2F(T)),\\p_{1.8}^{+}(T)&=Q(S_T>1.8F(T)).\end{aligned}',
25:r'\begin{aligned}\widehat V(T)&=\sum_{j=1}^{J}w_j\frac{u_{j-1}^2+u_{j-1}u_j+u_j^2}{3}-\widehat\mu^2,\\ V_{\mathrm{known}}(T)&=\sum_{\ell=1}^{L}\pi_\ell(e^{\sigma_\ell^2T}-1).\end{aligned}',
26:r'\begin{aligned}y_{\nu mi}(a)&=C_{mi}+\eta_{mi}(a)U_{\nu mi},\\ \eta_{mi}(a)&=\min\{a,C_{mi}/2\},\qquad U_{\nu mi}\sim\mathcal U[-1,1].\end{aligned}',
27:r'\begin{aligned}\overline\theta&=\frac{1}{n_{\mathrm{sim}}}\sum_{\nu=1}^{n_{\mathrm{sim}}}\widehat\theta_\nu,\qquad\widehat{\operatorname{Bias}}=\overline\theta-\theta,\\ s_\theta^2&=\frac{1}{n_{\mathrm{sim}}-1}\sum_{\nu=1}^{n_{\mathrm{sim}}}(\widehat\theta_\nu-\overline\theta)^2,\\ \operatorname{RMSE}_\theta^2&=\frac{1}{n_{\mathrm{sim}}}\sum_{\nu=1}^{n_{\mathrm{sim}}}(\widehat\theta_\nu-\theta)^2\\ &=\widehat{\operatorname{Bias}}^2+\frac{n_{\mathrm{sim}}-1}{n_{\mathrm{sim}}}s_\theta^2.\end{aligned}',
29:r'\begin{aligned}\overline\Delta&=\frac{1}{n_{\mathrm{sim}}}\sum_{\nu=1}^{n_{\mathrm{sim}}}\Delta_\nu,\\ \operatorname{MCSE}(\overline\Delta)&=\sqrt{\frac{\sum_{\nu=1}^{n_{\mathrm{sim}}}(\Delta_\nu-\overline\Delta)^2}{n_{\mathrm{sim}}(n_{\mathrm{sim}}-1)}}.\end{aligned}',
32:r'\begin{aligned}\widehat g_m^{(\varepsilon),\mathrm{ref}}(x)&=F_m\widehat f_m^{(\varepsilon)}(F_m x),\\ \widehat p_{1.8,m}^{(\varepsilon)}&=\int_{1.8}^{\infty}\widehat g_m^{(\varepsilon),\mathrm{ref}}(x)\,\mathrm dx.\end{aligned}',
33:r'\begin{aligned}C_m(K_{mi})&\geq L_{mi}^{(\varepsilon)}=D_m(\widetilde F_m^{(\varepsilon)}-K_{mi})_+,\\ v_{mi}^{(\varepsilon)}&=(L_{mi}^{(\varepsilon)}-y_{mi})_+.\end{aligned}',
34:r'\begin{aligned}\underline F_m&=\max_i\ell^F_{mi},\qquad\overline F_m=\min_i u^F_{mi},\\ \widehat F_m&=\frac{\underline F_m+\overline F_m}{2}\quad\text{when }0<\underline F_m\leq\overline F_m.\end{aligned}',
35:r'\begin{aligned}s_m^*&=\min\{s\geq0:\max_i(\ell^F_{mi}-s)\leq\min_i(u^F_{mi}+s)\}\\ &=\frac{(\underline F_m-\overline F_m)_+}{2}.\end{aligned}',
38:r'\begin{aligned}C_{mi}-P_{mi}&=H_m-D_mK_{mi},\\ b^-_{mi}&=C^{\mathrm{bid}}_{mi}-P^{\mathrm{ask}}_{mi},\\ b^+_{mi}&=C^{\mathrm{ask}}_{mi}-P^{\mathrm{bid}}_{mi}.\end{aligned}',
39:r'\begin{aligned}(\widehat H_m,\widehat D_m)&=\underset{H,D}{\arg\min}\;\frac{1}{N_m}\sum_{i=1}^{N_m}\left(\frac{H-DK_{mi}-y_{mi}}{K_{0m}}\right)^2,\\ \text{subject to}\quad &b^-_{mi}\leq H-DK_{mi}\leq b^+_{mi}\quad\text{for every }i,\\ &D\geq\varepsilon_{\mathrm{carry}},\quad H\geq\varepsilon_{\mathrm{carry}}K_{0m},\\ &\varepsilon_{\mathrm{carry}}=10^{-8},\qquad\widehat F_m=\widehat H_m/\widehat D_m.\end{aligned}',
41:r'\begin{aligned}\ell^C_{mi}&=\max\{B_{mi,C},B_{mi,P}+D_m(F_m-K_{mi})\},\\ u^C_{mi}&=\min\{A_{mi,C},A_{mi,P}+D_m(F_m-K_{mi})\}.\end{aligned}',
42:r'\begin{aligned}\delta_J^*&=\min_{\delta\geq0,\,W\in\mathcal P_J}\delta\\ \text{subject to}\quad &\ell^C_{mi}-\delta\leq(\mathbf A_m\mathbf w_m)_i\leq u^C_{mi}+\delta\\ &\text{for every }m,i.\end{aligned}',
43:r'\begin{aligned}\widehat W&=\arg\min_{W\in\mathcal F_J}\mathcal R_J(W),\\ \mathcal R_J(W)&=\frac{1}{J-2}\sum_{m=1}^{M}\sum_{j=2}^{J-1}\left(\frac{w_{m,j+1}-2w_{mj}+w_{m,j-1}}{\Delta u}\right)^2.\end{aligned}',
44:r'\begin{aligned}\underline p_J(\mathbf a)&=\min_{W\in\mathcal F_J}\mathbf a^{\mathsf T}\operatorname{vec}(W),\\ \overline p_J(\mathbf a)&=\max_{W\in\mathcal F_J}\mathbf a^{\mathsf T}\operatorname{vec}(W).\end{aligned}',
45:r'\begin{aligned}\omega_*&=\frac{T_*-T_a}{T_b-T_a},\qquad T_a\leq T_*\leq T_b,\\ \mathbf w_*&=(1-\omega_*)\mathbf w_a+\omega_*\mathbf w_b.\end{aligned}',
46:r'\begin{aligned}d_1(\widehat g^{(a)},\widehat g^{(b)})&=\int_{0.3}^{2.2}|\widehat g^{(a)}(x)-\widehat g^{(b)}(x)|\,dx\\ &=\sum_{j=1}^{J}|w_j^{(a)}-w_j^{(b)}|.\end{aligned}'
}
for k in (30,36):
    source=ASSETS['equations'][str(k)]['latex']
    left,right=source.split('\\qquad',1)
    BREAKS[k]='\\begin{aligned}'+left.rstrip(',\n ')+'\\\\'+right.strip()+'\\end{aligned}'

def clean_caption(s):
    s=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',s)
    return s.replace('Notebook 9\'s','the first empirical study\'s')

FIG4='Figure 4 introduces independent zero-mean quote perturbations with a standard deviation of 0.05 index points and seed 42. These errors need not preserve arbitrage constraints. Applying Equation (4) amplifies their effect through division by the squared strike spacing. This is an algebraic consequence of the stencil in Driscoll and Braun [7], rather than a calibrated model of market noise. The stress test deliberately shows the consequence of applying raw differences to inconsistent quotes; it is not a fitted market-noise model.'

headings=[]
source=(HERE/'manuscript.md').read_text()
for m in re.finditer(r'(?m)^(#{1,2}) (.+)$',source):
    key='h'+str(len(headings)+1)
    headings.append({'key':key,'title':m[2],'level':len(m[1])})

def expand(kind,num):
    n=int(num);a=ASSETS[{'EQ':'equations','FIG':'figures','TABLE':'tables'}[kind]][num]
    if kind=='EQ':
        # Single-column equation arrays avoid empty alignment cells that some
        # Word-compatible PDF exporters incorrectly render as error glyphs.
        latex=BREAKS.get(n,a['latex']).replace('{aligned}','{gathered}').replace('&','')
        return '\n\n$$\n'+latex+'\n$$\n\nEquation '+num+'. '+a['title']+'.\n\n'
    if kind=='FIG':
        landscape=n not in {1,2,3,4,6,7,8,9,10,14,28}
        w,h=Image.open(ROOT/a['path']).size
        width=min(9.85,4.9*w/h) if landscape else min(6.6,6.2*w/h)
        im=f'![{a["title"]}]({ROOT/a["path"]}){{width={width:.3f}in}}'
        inside=im+'\n\n'+clean_caption(a['caption'])
        if n==4:inside+='\n\n'+FIG4
        return f'\n\nFLOATSTART_{"L" if landscape else "P"}_FIG_{num}\n\n{inside}\n\nFLOATEND\n\n'
    cols=len(a['markdown'].splitlines()[0].split('|'))-2
    landscape=cols>=5 and n not in {10,11,13,17,18,20}
    table=a['markdown']
    if n==1:
        labels={'minimum_density':'Minimum density (per index point)','negative_density_count':'Negative density estimates','integrated_absolute_error':'Integrated absolute density error','first_moment':'Recovered mean (index points)','theoretical_forward':'Theoretical forward (index points)','forward_relative_error':'Relative error in the forward mean','max_repricing_error_index_points':'Maximum repricing error (index points)','mass':'Integrated probability mass'}
        for before,after in labels.items():table=table.replace('| '+before+' |','| '+after+' |')
    return f'\n\nFLOATSTART_{"L" if landscape else "P"}_TABLE_{num}\n\n{table}\n\n{clean_caption(a["caption"])}\n\nFLOATEND\n\n'

source=source.replace('{{FIG4EXPLANATION}}','')
source=re.sub(r'\{\{(EQ|FIG|TABLE):(\d+)\}\}',lambda m:expand(m[1],m[2]),source)
refs=(ROOT/'research/references.md').read_text().split('\n## Relationship-to-source guide',1)[0]
refs=refs[refs.index('[1]'):].strip()
refs=refs.replace('in Milestone 6','for the mixture study')
source=source.replace('{{REFERENCES}}',refs)
# Adjacent landscape assets share the same page orientation without an empty
# portrait section between them. Each retains a separate printable page.
source=re.sub(r'FLOATEND\s+FLOATSTART_L_', 'FLOATNEXT_L_',source)
front=f'COVERTITLE\n\n{TITLE}\n\n{SUBTITLE}\n\nJoel Cerraga\n\nQuantitative research project\n\nSeptember 2026\n\nFinal assembled paper · Version 1.0\n\nCOVEREND\n\n# Abstract\n\n{ABSTRACT}\n\n# Table of contents\n\n'
frontheads=['Abstract','List of figures','List of tables','List of equations','List of abbreviations','List of mathematical symbols']
for title in frontheads:
    key='front_'+title.lower().replace(' ','_')
    front+=f'NAV|{key}|0|{title}\n\n'
for h in headings:front+=f'NAV|{h["key"]}|{h["level"]}|{h["title"]}\n\n'
for kind,tag in [('figures','Figure'),('tables','Table'),('equations','Equation')]:
    front+='# List of '+kind+'\n\n'
    for n,a in ASSETS[kind].items():front+=f'NAV|{tag.lower()}{n}|1|{tag} {n}. {a["title"]}\n\n'
abbr=ASSETS['registers']['abbreviations']
front+='# List of abbreviations\n\n'+'\n'.join(abbr[:2]+sorted(abbr[2:],key=str.casefold))+'\n\n'
front+='# List of mathematical symbols\n\n'
front+='Symbols are defined in the context of their equation. Reused letters with different subscripts or roles are distinguished below.\n\n'
front+='\n'.join(ASSETS['registers']['mathematical symbols'])+'\n\nBODYBEGIN\n\n'
(HERE/'expanded.md').write_text(front+source)
subprocess.run(['pandoc',str(HERE/'expanded.md'),'-f','markdown+tex_math_dollars-implicit_figures','-t','docx','-o',str(HERE/'draft.docx')],check=True)
doc=Document(HERE/'draft.docx')

def xml(tag,**attrs):
    el=OxmlElement(tag)
    for k,v in attrs.items():el.set(qn('w:'+k),str(v))
    return el

styles=doc.styles
for st in styles:
    if re.match(r'Heading [1-9]$',st.name):st.name=st.name.lower()
for name in ['Front Heading','Nav Entry','Equation Caption','Figure Caption','Table Caption','Reference','Cover Title','Cover Subtitle','Cover Detail']:
    if name not in styles:styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
normal=styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(8)
normal.paragraph_format.widow_control=True
for name in ['Body Text','First Paragraph']:
    styles[name].paragraph_format.line_spacing=1.15
    styles[name].paragraph_format.space_after=Pt(8)
    styles[name].paragraph_format.first_line_indent=Pt(0)
for st in styles:
    if st.type==WD_STYLE_TYPE.PARAGRAPH:
        st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
        if st.name not in ['Normal'] and st.base_style is None:st.base_style=normal
for name,size in [('Title',22),('Heading 1',16),('Heading 2',13),('Heading 3',12),('Front Heading',16)]:
    st=styles[name];st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(16);st.paragraph_format.space_after=Pt(10)
    st.paragraph_format.keep_with_next=True
    if name in ['Heading 1','Front Heading']:st.paragraph_format.page_break_before=True
for name in ['Figure Caption','Table Caption','Equation Caption']:
    st=styles[name];st.font.size=Pt(10);st.font.italic=True
    st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_before=Pt(5);st.paragraph_format.space_after=Pt(12)
    st.paragraph_format.keep_together=True
styles['Equation Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Nav Entry'].font.size=Pt(10.5)
styles['Nav Entry'].paragraph_format.space_after=Pt(5)
styles['Nav Entry'].paragraph_format.line_spacing=1.05
styles['Reference'].font.size=Pt(10.5)
styles['Reference'].paragraph_format.line_spacing=1.1
styles['Reference'].paragraph_format.space_after=Pt(10)
styles['Reference'].paragraph_format.left_indent=Inches(.30)
styles['Reference'].paragraph_format.first_line_indent=Inches(-.30)
for name,size in [('Cover Title',24),('Cover Subtitle',15),('Cover Detail',12)]:
    styles[name].font.size=Pt(size)
    styles[name].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    styles[name].paragraph_format.space_after=Pt(18)
    styles[name].paragraph_format.keep_with_next=True
styles['Cover Title'].font.bold=True
styles['Cover Title'].paragraph_format.space_before=Pt(100)
styles['Cover Subtitle'].paragraph_format.space_after=Pt(70)
styles['Title'].font.size=Pt(24)
styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Title'].paragraph_format.space_before=Pt(100)
styles['Title'].paragraph_format.space_after=Pt(18)
styles['Hyperlink'].font.color.rgb=RGBColor.from_string('1F4E79')

# Add native multilevel chapter numbering rather than typed heading numbers.
numbering=doc.part.numbering_part.element
absnum=xml('w:abstractNum',abstractNumId=70)
absnum.append(xml('w:multiLevelType',val='multilevel'))
for lev in range(2):
    lvl=xml('w:lvl',ilvl=lev);lvl.append(xml('w:start',val=1));lvl.append(xml('w:numFmt',val='decimal'))
    lvl.append(xml('w:lvlText',val='%1' if lev==0 else '%1.%2'))
    lvl.append(xml('w:suff',val='space'));lvl.append(xml('w:lvlJc',val='left'))
    absnum.append(lvl)
numbering.append(absnum)
num=xml('w:num',numId=70);num.append(xml('w:abstractNumId',val=70));numbering.append(num)

bookmark_id=500
def bookmark(p,name):
    global bookmark_id
    bookmark_id+=1
    start=xml('w:bookmarkStart',id=bookmark_id,name=name);end=xml('w:bookmarkEnd',id=bookmark_id)
    p.insert(0,start);p.append(end)

def addfield(p,instruction,text='1'):
    for typ,value in [('begin',None),('instr',instruction),('separate',None),('text',text),('end',None)]:
        r=OxmlElement('w:r')
        if typ in ['begin','separate','end']:r.append(xml('w:fldChar',fldCharType=typ))
        else:
            el=OxmlElement('w:instrText' if typ=='instr' else 'w:t');el.text=value;r.append(el)
        p.append(r)

def section(mode='P',fmt='decimal',start=None,cover=False):
    s=xml('w:sectPr')
    s.append(xml('w:type',val='nextPage'))
    s.append(xml('w:pgSz',w=16838 if mode=='L' else 11906,h=11906 if mode=='L' else 16838,orient='landscape' if mode=='L' else 'portrait'))
    s.append(xml('w:pgMar',top=1008,bottom=1008,left=1152,right=1152,header=432,footer=432,gutter=0))
    pgn=xml('w:pgNumType',fmt=fmt)
    if start is not None:pgn.set(qn('w:start'),str(start))
    s.append(pgn)
    if cover:s.append(xml('w:titlePg'))
    return s

def end_section(p,mode='P',fmt='decimal',start=None,cover=False):
    p.clear();pp=xml('w:pPr');pp.append(section(mode,fmt,start,cover));p.append(pp)
    pp.append(xml('w:spacing',before=0,after=0,line=1,lineRule='exact'))
    r=xml('w:r');rp=xml('w:rPr');rp.append(xml('w:sz',val=2));r.append(rp);p.append(r)

mode='P';float_kind=None;current_table=None;infront=True;cover=True;inrefs=False;hindex=0;float_elements=[];last_el=None
body=doc.element.body
for old in list(body):
    if old.tag==qn('w:sectPr'):body.remove(old)
elements=list(body)
for el in elements:
    if el.tag==qn('w:tbl'):
        t=Table(el,doc);cols=len(t.columns);width=9.95 if mode=='L' else 6.65
        t.autofit=False
        if infront and cols==3:fractions=[.23,.51,.26]
        elif infront and cols==2:fractions=[.18,.82]
        elif cols==2:fractions=[.50,.50]
        elif cols==3:fractions=[.27,.24,.49] if current_table==5 else [.23,.32,.45]
        elif cols==4:fractions=[.27,.25,.25,.23]
        elif cols==5:fractions=[.22,.22,.20,.18,.18]
        elif cols==6:fractions=[.21,.16,.16,.16,.16,.15]
        else:fractions=[1/cols]*cols
        for j,c in enumerate(t.columns):c.width=Inches(width*fractions[j])
        borders=xml('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:borders.append(xml('w:'+side,val='single',sz=4,color='D9D9D9'))
        t._tbl.tblPr.append(borders)
        margin=xml('w:tblCellMar')
        for side in ['top','bottom']:margin.append(xml('w:'+side,w=75,type='dxa'))
        for side in ['left','right']:margin.append(xml('w:'+side,w=85,type='dxa'))
        t._tbl.tblPr.append(margin)
        for ri,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();trpr.append(xml('w:cantSplit'))
            if ri==0:trpr.append(xml('w:tblHeader'))
            for ci,c in enumerate(row.cells):
                c.width=Inches(width*fractions[ci]);c._tc.get_or_add_tcPr().append(xml('w:vAlign',val='center'))
                if ri==0:c._tc.get_or_add_tcPr().append(xml('w:shd',fill='E9EEF3'))
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.05
                    p.paragraph_format.keep_with_next=(ri==0 or (not infront and (len(t.rows)<=12 or ri==len(t.rows)-1)))
                    p.paragraph_format.keep_together=True
                    txt=p.text.strip()
                    if ri>0 and re.match(r'^[−+\-]?\d',txt) and len(txt)<30:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                    for r in p.runs:r.font.size=Pt(10 if infront or mode=='L' else 10.5);r.bold=(ri==0)
        last_el=el;continue
    if el.tag!=qn('w:p'):continue
    p=Paragraph(el,doc);txt=p.text
    if txt=='COVERTITLE':body.remove(el);continue
    if cover and txt=='COVEREND':
        end_section(el,cover=True);cover=False;continue
    if cover:
        p.style='Title' if txt==TITLE else 'Cover Subtitle' if txt==SUBTITLE else 'Cover Detail'
        last_el=el;continue
    if txt=='BODYBEGIN':end_section(el,fmt='lowerRoman',start=1);infront=False;continue
    if txt.startswith('NAV|'):
        _,key,level,label=txt.split('|',3)
        p.clear();p.style='Nav Entry'
        level=int(level)
        p.paragraph_format.left_indent=Inches(.18 if level==2 else 0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.65-(.18 if level==2 else 0)),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        link=xml('w:hyperlink',anchor=key,history=1)
        for textval in [label,'\t',str(PAGES.get(key,'1'))]:
            r=xml('w:r');rp=xml('w:rPr');rp.append(xml('w:color',val='000000'))
            if level==1 and key.startswith('h'):rp.append(xml('w:b'))
            r.append(rp)
            if textval=='\t':r.append(xml('w:tab'))
            else:tx=xml('w:t');tx.text=textval;r.append(tx)
            link.append(r)
        p._p.append(link);continue
    if txt.startswith('FLOATNEXT_'):
        _,newmode,kind,number=txt.split('_');current_table=int(number) if kind=='TABLE' else None;float_kind=kind
        end_section(el,mode);mode='L'
        continue
    if txt.startswith('FLOATSTART_'):
        _,newmode,kind,number=txt.split('_');current_table=int(number) if kind=='TABLE' else None;float_kind=kind
        if newmode=='L':end_section(el,mode);mode='L'
        else:body.remove(el)
        continue
    if txt=='FLOATEND':
        if mode=='L':end_section(el,'L');mode='P'
        else:body.remove(el)
        current_table=None;float_kind=None;continue
    if p.style.name.startswith('Heading'):
        if infront:
            p.style='Front Heading';bookmark(el,'front_'+txt.lower().replace(' ','_'))
            if txt=='Abstract':p.paragraph_format.page_break_before=False
        else:
            h=headings[hindex];hindex+=1;bookmark(el,h['key'])
            # Pandoc creates its own bookmark too; our stable keys drive navigation.
            if re.match(r'^\d+(?:\.\d+)? ',txt):
                number,rest=txt.split(' ',1)
                for child in list(el):
                    if child.tag not in [qn('w:pPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')]:el.remove(child)
                p.add_run(rest)
                pp=p._p.get_or_add_pPr();np=xml('w:numPr');np.append(xml('w:ilvl',val=h['level']-1));np.append(xml('w:numId',val=70));pp.append(np)
            if hindex==1:p.paragraph_format.page_break_before=False
            inrefs=txt=='References' or (inrefs and not txt.startswith('Appendix'))
        last_el=el;continue
    cap=re.match(r'^(Equation|Figure|Table) (\d+)\. ',txt)
    if cap:
        kind,n=cap.groups();p.style=kind+' Caption';bookmark(el,kind.lower()+n)
        if kind in ['Equation','Figure'] and last_el is not None and last_el.tag==qn('w:p'):
            Paragraph(last_el,doc).paragraph_format.keep_with_next=True
        if kind=='Figure' and n=='4':p.paragraph_format.keep_with_next=True
    elif el.xpath('.//m:oMathPara'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(2)
    elif el.xpath('.//w:drawing'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(2)
    elif inrefs and re.match(r'^\[\d+\]',txt):p.style='Reference'
    elif txt==FIG4:p.paragraph_format.keep_together=True
    last_el=el
body.append(section('P'))

# First body section restarts decimal numbering; subsequent landscape sections continue.
sects=doc.sections
for i,s in enumerate(sects):
    if i==2:s._sectPr.find(qn('w:pgNumType')).set(qn('w:start'),'1')
    s.footer_distance=Inches(.3);s.header_distance=Inches(.3)
    if i in [0,1,2]:
        s.footer.is_linked_to_previous=False
        fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        if i!=0:addfield(fp._p,' PAGE ')
        for r in fp.runs:r.font.size=Pt(10);r.font.name='Arial'
        s.header.is_linked_to_previous=False
        hp=s.header.paragraphs[0]
        if i==2:
            hp.text='OPTION-IMPLIED DENSITIES  |  JOEL CERRAGA'
            hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            for r in hp.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('555555')
    # Inherited headers/footers across page orientation changes retain correct numbering.

doc.core_properties.title=TITLE
doc.core_properties.subject=SUBTITLE
doc.core_properties.author='Joel Cerraga'
doc.core_properties.keywords='SPX; risk-neutral density; constrained estimation; option prices'
doc.core_properties.comments='Final paper assembled from the completed ten-notebook research project.'
settings=doc.settings.element
settings.append(xml('w:updateFields',val='false'))
mp=settings.find(qn('m:mathPr'))
if mp is None:mp=OxmlElement('m:mathPr');settings.append(mp)
mf=OxmlElement('m:mathFont');mf.set(qn('m:val'),'Cambria Math');mp.append(mf)
out=HERE/(FILENAME+'.docx');doc.save(out)
(HERE/'navigation.json').write_text(json.dumps({'headings':headings,'front':frontheads,'title':TITLE},indent=2))
print(out)
print('Native math objects:',len(doc.element.xpath('//m:oMath')),'Sections:',len(doc.sections))
